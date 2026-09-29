# Regenerates docs/ROLE_GUIDE.xlsx: python scripts/build-role-guide-xlsx.py docs/ROLE_GUIDE.xlsx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="14532D")
YES_FILL = PatternFill("solid", fgColor="BBF7D0")
NO_FILL = PatternFill("solid", fgColor="F3F4F6")
BAND_FILL = PatternFill("solid", fgColor="ECFDF5")
thin = Side(style="thin", color="D1D5DB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def head(ws, row, values):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 34


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color="14532D")
    ws["A2"] = sub
    ws["A2"].font = Font(name=FONT, italic=True, size=9, color="6B7280")


wb = Workbook()

# --- Sheet 1: permission matrix -------------------------------------------
ws = wb.active
ws.title = "Yetki Matrisi"
roles = ["Firma Admini", "Üretim Sorumlusu", "Kalite Sorumlusu", "Depocu (Operator)",
         "Firma Kullanıcısı", "Bölge Müdürü", "Salt Okuma"]
title(ws, "Kim neyi değiştirebilir?",
      "✓ = yapabilir, – = yapamaz. Kaynak: types/roles.ts. Son gözden geçirme: 2026-09-29.")
head(ws, 4, ["İşlem"] + roles)
A, P, Q, O, U, R, V = range(1, 8)
rows = [
    ("Hammadde, ambalaj, ürün, reçete, tedarikçi ekleme/düzenleme", [A, P, U]),
    ("Üretim emri açma, başlatma, tamamlama", [A, P, O, U]),
    ("Kalite kontrol açma / imzalama", [A, Q, U]),
    ("Stok: lot, mal kabul, barkod okutma, LTD deposuna sayarak alma", [A, P, O, U]),
    ("Dosya ekleme (analiz sertifikası, MSDS, rapor)", [A, P, Q, U]),
    ("Sevkiyat hazırlama ve gönderme", [A, P, O, U]),
    ("Eczane siparişi girme / yönetme", [A, P, O, U, R]),
    ("Pazaryeri: liste, stok gönderme, fiyat onayı", [A, P, O, U]),
    ("Ajan Kurulu (yapay zekâ)", [A, P, U]),
    ("Kullanıcı davet etme, rol verme", [A]),
    ("Pazaryeri API bağlantı bilgileri", [A]),
    ("Depo ve raf tanımlama (Ayarlar)", [A]),
    ("Üretim, reçete, maliyet ekranlarını görme", [A, P, Q, U, V]),
    ("Sadece bitmiş ürün stoğunu görme (hammadde/fabrika görmez)", [O]),
    ("Sadece sipariş, ürün, stok, üretim (salt okuma), müşteri görme", [R]),
]
for r, (name, allowed) in enumerate(rows, 5):
    c = ws.cell(row=r, column=1, value=name)
    c.font = Font(name=FONT, size=10)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = BORDER
    for k in range(1, 8):
        cell = ws.cell(row=r, column=k + 1, value="✓" if k in allowed else "–")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        cell.fill = YES_FILL if k in allowed else NO_FILL
        cell.font = Font(name=FONT, bold=k in allowed, size=11,
                         color="166534" if k in allowed else "9CA3AF")
    ws.row_dimensions[r].height = 30
last = 4 + len(rows)
ws.cell(row=last + 2, column=1,
        value="Eczane alıcısı ERP'ye giremez; yalnızca /portal'dan katalog görür ve sipariş verir. "
              "Platform admini yalnızca /superadmin'de firma, paket ve kullanıcı yönetir; firma verisine dokunmaz."
        ).font = Font(name=FONT, size=9, italic=True, color="6B7280")
ws.column_dimensions["A"].width = 58
for i in range(2, 9):
    ws.column_dimensions[get_column_letter(i)].width = 14
ws.freeze_panes = "B5"

# --- Sheet 2: how-to -------------------------------------------------------
ws2 = wb.create_sheet("Nasıl Yapılır")
title(ws2, "Hangi ekrandan, nasıl yapılır?",
      "Filtreyi kullanarak Rol sütunundan kendi rolünüzü seçin.")
head(ws2, 4, ["Rol", "İşlem", "Nereden (menü)", "Nasıl"])
D, Y, K = "Depocu", "Yönetici*", "Kalite"
howto = [
    (D, "Partiyi LTD deposuna sayarak almak",
     "Parti kartındaki QR (telefon kamerası) veya Depo Hareketleri → Barkod Tara → Kamerayla Tara",
     "1) Lotu açın 2) 'Sayılan miktar'a sayıyı yazın 3) 'NeuPharma LTD.ŞTİ.' butonuna basın"),
    (D, "Sayım sistemle aynıysa", "Aynı ekran", "Lot hemen alınır; '… stoğuna alındı' yazar"),
    (D, "Sayım farklıysa", "Aynı ekran",
     "Sarı kutu farkı gösterir → nedenini yazın → '… olarak stoğa al'. Sayılan miktar stok olur, fark düzeltme olarak kaydolur"),
    (D, "Yanlış yazdıysa", "Aynı ekran", "'Vazgeç' butonu"),
    (D, "Kamera yoksa", "Barkod Tara", "Lot numarasını kutuya yazıp 'Bul'"),
    (D, "Aynı depoda rafa yerleştirmek", "Barkod Tara", "Lotu okutun → raf etiketini okutun veya listeden seçin (sayım gerekmez)"),
    (D, "Raftakileri görmek", "Barkod Tara", "Raf etiketini okutun"),
    (D, "Sipariş hazırlamak", "Eczane Siparişleri veya Siparişler → Yeni", "Siparişi açın → sevkiyata çevirin → serbest lotları seçin"),
    (D, "Göndermek", "Siparişler", "Sevkiyatı 'Gönder' (gönderilen değiştirilemez)"),
    (D, "Pazaryeri fiyat onayı", "Pazaryeri", "Öneriyi onaylayın veya reddedin"),
    (D, "Mevcut stoğu girmek", "Stok Girişi (Barkod)", "Ürün barkodunu okutun, adedi girin"),
    (Y, "Reçete oluşturmak", "Reçeteler → Yeni",
     "YM reçetesine yalnızca hammadde; mamül reçetesine yalnızca YM + ambalaj. Kullanmadan önce yayınlayın"),
    (Y, "Üretim emri açmak", "Üretim → Yeni Üretim Emri", "Yayında reçete + hedef miktar (mamülde KUTU)"),
    (Y, "Üretimi başlatmak", "Emir sayfası",
     "Planla → Üretime Al. Depoda tek YM partisi varsa parti no ondan önerilir; parti no ürün bazında benzersiz"),
    (Y, "Üretimi tamamlamak", "Emir sayfası",
     "Tamamla: tüketilen lotlar, çıkış lot no, SKT, konum. Çıkış lotu karantinada açılır"),
    (Y, "Parti kartı basmak", "Tamamlanmış emir", "Parti Kartı (PDF) → PDF Kaydet / Yazdır"),
    (Y, "Emri düzeltmek / silmek", "Emir sayfası", "Düzenle (tamamlanana kadar); Sil (Silme Protokolü'ne kaydolur)"),
    (Y, "Lotu yönetmek", "Lotlar", "Düzenle, yanlış lotu sil, Etiket Yazdır (QR)"),
    (Y, "Mal kabul / elle stok hareketi", "Depo Hareketleri", "Mal Kabul, Stok Hareketi"),
    (Y, "Maliyetsiz lotlara maliyet girmek", "Stok", "Stok Değerleme"),
    (Y, "Karantinadan çıkarmak (satılabilir yapmak)", "Lotlar → lot", "Durumu 'Serbest' yapın"),
    (Y, "Yanlış stok hareketini geri almak", "Lot / hareket", "Storno"),
    (Y, "İhtiyaç hesabı", "MRP (İhtiyaç Planlama)", "Açık siparişlerden hesaplanır"),
    (Y, "Satınalma siparişi ve mal kabul", "Satınalma Siparişleri", "Tedarikçiye gönder, mal kabulü kaydet"),
    (Y, "Fatura / irsaliye", "Fatura / İrsaliye", "Yeni"),
    (Y, "Aylık gider girmek", "Genel Giderler", "Maaş, elektrik, su, kira vb. girin"),
    (Y, "Müşteri / tedarikçi hesabı", "Cari Hesaplar", "Ödeme için dekont no, tarih, tutar zorunlu"),
    (Y, "Maliyet ve karlılık raporu", "Raporlar", "Maliyet, karlılık, hammadde maliyetleri"),
    (Y, "Eczane siparişi işlemek", "Eczane Siparişleri", "Siparişi açın, sevkiyata çevirin"),
    (Y, "Portal kataloğu ve alıcı girişi", "Portal Kataloğu, Eczane Hesapları", "Ürünleri listeleyin, alıcı girişi tanımlayın"),
    (Y, "Ajan Kurulu çalıştırmak", "Ajan Kurulu", "Tartışmayı çalıştırın; önerilen aksiyonlar onaydan sonra uygulanır"),
    ("Firma Admini", "Kullanıcı davet etmek, rol vermek", "Kullanıcılar", "Davet et, rol seç"),
    ("Firma Admini", "Pazaryeri API bilgileri", "Ayarlar → Pazaryeri Bağlantıları", "Bilgileri girin"),
    ("Firma Admini", "Depo ve raf tanımı", "Ayarlar", "Yeni depo / raf ekleyin"),
    (K, "Kalite kontrol açmak", "Kalite Kontrol → Yeni", "Lot veya parti seçin"),
    (K, "İmzalamak / iptal etmek", "Kontrol sayfası", "Sonucu girip imzalayın (başarısız kontrol lotu serbest bırakmayı engeller)"),
    (K, "Dosya eklemek", "Lot / kontrol sayfası", "Analiz sertifikası, MSDS, rapor yükleyin"),
    ("Bölge Müdürü", "Eczane adına sipariş girmek", "Eczane Siparişleri", "Yeni sipariş; eczane ve ürünleri seçin"),
    ("Eczane alıcısı", "Katalog görmek, sipariş vermek", "/portal", "Kendi girişiyle"),
    ("Platform admini", "Firma, paket, koltuk limiti, kullanıcı daveti, denetim kaydı", "/superadmin",
     "Firmalar / Abonelikler / Kullanıcılar / Audit"),
]
for r, row in enumerate(howto, 5):
    for k, v in enumerate(row, 1):
        c = ws2.cell(row=r, column=k, value=v)
        c.font = Font(name=FONT, size=10, bold=(k == 1))
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
        if r % 2 == 0:
            c.fill = BAND_FILL
    ws2.row_dimensions[r].height = 44
ws2.cell(row=5 + len(howto) + 1, column=1,
         value="* Yönetici = Firma Admini, Üretim Sorumlusu, Firma Kullanıcısı. Kullanıcı yönetimi yalnızca Firma Admini'ndedir."
         ).font = Font(name=FONT, size=9, italic=True, color="6B7280")
for col, w in zip("ABCD", [16, 40, 46, 70]):
    ws2.column_dimensions[col].width = w
ws2.auto_filter.ref = f"A4:D{4 + len(howto)}"
ws2.freeze_panes = "A5"

# --- Sheet 3: flow ----------------------------------------------------------
ws3 = wb.create_sheet("Akış")
title(ws3, "Fabrikadan müşteriye akış", "Sıra numarasına göre.")
head(ws3, 4, ["Adım", "Kim", "Ne olur"])
flow = [
    (1, "Yönetici", "Hammadde gelir (Mal Kabul / satınalma kabulü) → karantina → kalite → serbest"),
    (2, "Yönetici", "YM emri → tamamla → YM lotu (parti no örn. 2608006)"),
    (3, "Yönetici", "Aynı parti numarasıyla mamül emri → tamamla → mamül lotu ANA depoda karantinada"),
    (4, "Yönetici", "Parti Kartı (PDF) yazdırılır, ürünün yanına konur"),
    (5, "Depocu", "QR okutur, sayar, lotu LTD deposuna alır"),
    (6, "Yönetici / Kalite", "Lot serbest bırakılır → satılabilir stok, pazaryeri ve portal stoğu"),
    (7, "Depocu", "Sipariş gelir (portal, pazaryeri, elle) → sevkiyat → gönder"),
]
for r, row in enumerate(flow, 5):
    for k, v in enumerate(row, 1):
        c = ws3.cell(row=r, column=k, value=v)
        c.font = Font(name=FONT, size=10, bold=(k == 1))
        c.alignment = Alignment(wrap_text=True, vertical="center",
                                horizontal="center" if k == 1 else "left")
        c.border = BORDER
        if r % 2 == 0:
            c.fill = BAND_FILL
    ws3.row_dimensions[r].height = 30
ws3.cell(row=13, column=1, value="Açık sorular (sahibi karar verecek)").font = Font(name=FONT, bold=True, size=10)
ws3.cell(row=14, column=1, value="• Depo lotu aldığında lot otomatik serbest kalsın mı, yoksa kalite kararı olarak mı kalsın?").font = Font(name=FONT, size=10)
ws3.cell(row=15, column=1, value="• Sayım farkında kabul yönetici onayı beklesin mi?").font = Font(name=FONT, size=10)
for col, w in zip("ABC", [8, 22, 100]):
    ws3.column_dimensions[col].width = w

import sys
wb.save(sys.argv[1])
print("saved", sys.argv[1])
